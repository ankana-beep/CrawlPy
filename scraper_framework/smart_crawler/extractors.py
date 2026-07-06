"""Generic extractors for unknown unauthenticated websites."""

from __future__ import annotations

import csv
import io
import json
import re
import xml.etree.ElementTree as ET
from typing import Any, Dict, Iterable, List, Optional

from bs4 import BeautifulSoup

from scraper_framework.smart_crawler.models import ExtractedPage, FetchResult
from scraper_framework.smart_crawler.url_tools import (
    looks_like_api,
    looks_like_feed,
    looks_like_file,
    normalize_url,
)

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?:\+?\d[\d\s().-]{7,}\d)")


class GenericExtractor:
    def extract_html(self, html: str, base_url: str, *, include_raw: bool = True) -> ExtractedPage:
        soup = BeautifulSoup(html, "lxml")
        links = _extract_links(soup, base_url)
        feeds = _extract_feeds(soup, base_url, links)
        sitemaps = _extract_sitemaps(soup, base_url, links)
        files = [url for url in links if looks_like_file(url)]
        api_candidates = [url for url in links if looks_like_api(url)]
        embedded_json = _extract_embedded_json(html)
        api_candidates.extend(_extract_api_urls_from_text(html, base_url))

        for element in soup(["script", "style", "noscript"]):
            element.extract()

        title = soup.title.string.strip() if soup.title and soup.title.string else None
        text = soup.get_text(separator=" ", strip=True)

        return ExtractedPage(
            url=base_url,
            title=title,
            text=text,
            meta=_extract_meta(soup),
            links=links,
            feeds=feeds,
            sitemaps=sitemaps,
            files=files,
            api_candidates=_dedupe(api_candidates),
            embedded_json=embedded_json,
            tables=_extract_tables(soup),
            forms=_extract_forms(soup, base_url),
            emails=sorted(set(EMAIL_RE.findall(text))),
            phones=sorted(set(match.strip() for match in PHONE_RE.findall(text))),
            raw_html=html if include_raw else None,
        )

    def extract_fetch_result(self, result: FetchResult) -> Dict[str, Any]:
        content_type = result.content_type.lower()
        url = result.final_url or result.url
        body_text = result.body_text or ""
        if "json" in content_type or url.lower().endswith(".json"):
            return {"kind": "json", "data": _safe_json(body_text)}
        if "xml" in content_type or url.lower().endswith((".xml", ".rss", ".atom")):
            return {"kind": "xml", "data": _parse_xml(body_text), "discovered_urls": _urls_from_xml(body_text)}
        if "csv" in content_type or url.lower().endswith(".csv"):
            return {"kind": "csv", "data": _parse_csv(body_text)}
        if "pdf" in content_type or url.lower().endswith(".pdf"):
            return {"kind": "pdf", "data": _parse_pdf(result.body_bytes or b"")}
        if url.lower().endswith((".xls", ".xlsx")):
            return {"kind": "excel", "data": _parse_excel(result.body_bytes or b"")}
        return {"kind": "binary" if result.body_text is None else "text", "data": body_text}


def _extract_links(soup: BeautifulSoup, base_url: str) -> List[str]:
    urls: List[str] = []
    attrs = [
        ("a", "href"),
        ("link", "href"),
        ("script", "src"),
        ("img", "src"),
        ("iframe", "src"),
        ("source", "src"),
    ]
    for tag_name, attr in attrs:
        for node in soup.find_all(tag_name):
            url = normalize_url(node.get(attr, ""), base_url)
            if url:
                urls.append(url)
    return _dedupe(urls)


def _extract_feeds(soup: BeautifulSoup, base_url: str, links: Iterable[str]) -> List[str]:
    feed_urls = []
    for node in soup.find_all("link"):
        rel = " ".join(node.get("rel") or []).lower()
        type_attr = (node.get("type") or "").lower()
        if "alternate" in rel and ("rss" in type_attr or "atom" in type_attr):
            url = normalize_url(node.get("href", ""), base_url)
            if url:
                feed_urls.append(url)
    feed_urls.extend(url for url in links if looks_like_feed(url))
    return _dedupe(feed_urls)


def _extract_sitemaps(soup: BeautifulSoup, base_url: str, links: Iterable[str]) -> List[str]:
    discovered = [url for url in links if "sitemap" in url.lower()]
    root = normalize_url("/sitemap.xml", base_url)
    if root:
        discovered.append(root)
    return _dedupe(discovered)


def _extract_meta(soup: BeautifulSoup) -> Dict[str, Any]:
    meta: Dict[str, Any] = {}
    for node in soup.find_all("meta"):
        key = node.get("name") or node.get("property") or node.get("itemprop")
        value = node.get("content")
        if key and value:
            meta[key] = value
    return meta


def _extract_embedded_json(html: str) -> List[Dict[str, Any]]:
    soup = BeautifulSoup(html, "lxml")
    out: List[Dict[str, Any]] = []
    for script in soup.find_all("script"):
        script_type = (script.get("type") or "").lower()
        text = script.string or script.get_text() or ""
        if not text.strip():
            continue
        if "json" in script_type or "ld+json" in script_type:
            parsed = _safe_json(text)
            out.append({"type": script_type or "json", "data": parsed})
            continue
        for candidate in _json_object_candidates(text):
            parsed = _safe_json(candidate)
            if parsed is not None:
                out.append({"type": "script_candidate", "data": parsed})
    return out


def _json_object_candidates(text: str) -> List[str]:
    candidates: List[str] = []
    for pattern in (r"window\.__INITIAL_STATE__\s*=\s*({.*?});", r"__NEXT_DATA__\s*=\s*({.*?});"):
        candidates.extend(re.findall(pattern, text, flags=re.DOTALL))
    return candidates[:20]


def _extract_api_urls_from_text(text: str, base_url: str) -> List[str]:
    raw_urls = re.findall(r"""["']([^"']*(?:/api/|/graphql|/wp-json/)[^"']*)["']""", text, flags=re.I)
    urls = []
    for raw in raw_urls:
        url = normalize_url(raw, base_url)
        if url:
            urls.append(url)
    return _dedupe(urls)


def _extract_tables(soup: BeautifulSoup) -> List[Dict[str, Any]]:
    tables = []
    for table in soup.find_all("table")[:20]:
        rows = []
        for tr in table.find_all("tr"):
            cells = [cell.get_text(" ", strip=True) for cell in tr.find_all(["th", "td"])]
            if cells:
                rows.append(cells)
        if rows:
            tables.append({"rows": rows})
    return tables


def _extract_forms(soup: BeautifulSoup, base_url: str) -> List[Dict[str, Any]]:
    forms = []
    for form in soup.find_all("form"):
        fields = []
        for field in form.find_all(["input", "select", "textarea", "button"]):
            fields.append(
                {
                    "tag": field.name,
                    "name": field.get("name"),
                    "type": field.get("type"),
                    "value": field.get("value"),
                    "text": field.get_text(" ", strip=True),
                }
            )
        forms.append(
            {
                "method": (form.get("method") or "GET").upper(),
                "action": normalize_url(form.get("action", ""), base_url),
                "fields": fields,
            }
        )
    return forms


def _urls_from_xml(text: str) -> List[str]:
    urls = re.findall(r"<loc>\s*([^<]+)\s*</loc>", text, flags=re.I)
    urls.extend(re.findall(r"<link>\s*([^<]+)\s*</link>", text, flags=re.I))
    return _dedupe(urls)


def _parse_xml(text: str) -> Dict[str, Any]:
    try:
        root = ET.fromstring(text.encode("utf-8"))
        return {"root": root.tag, "urls": _urls_from_xml(text)}
    except Exception as exc:
        return {"error": str(exc), "raw": text[:5000]}


def _parse_csv(text: str) -> Dict[str, Any]:
    reader = csv.DictReader(io.StringIO(text))
    rows = list(reader)
    if rows:
        return {"rows": rows[:500], "truncated": len(rows) > 500}
    simple_rows = list(csv.reader(io.StringIO(text)))
    return {"rows": simple_rows[:500], "truncated": len(simple_rows) > 500}


def _parse_pdf(data: bytes) -> Dict[str, Any]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return {"available": False, "reason": "Install pypdf to extract PDF text.", "size_bytes": len(data)}

    try:
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages[:100])
        return {"pages": len(reader.pages), "text": text}
    except Exception as exc:
        return {"error": str(exc), "size_bytes": len(data)}


def _parse_excel(data: bytes) -> Dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"available": False, "reason": "Install openpyxl to extract Excel rows.", "size_bytes": len(data)}

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheets: Dict[str, Any] = {}
        for sheet in workbook.worksheets:
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= 500:
                    break
                rows.append(list(row))
            sheets[sheet.title] = rows
        return {"sheets": sheets}
    except Exception as exc:
        return {"error": str(exc), "size_bytes": len(data)}


def _safe_json(text: str) -> Any:
    try:
        return json.loads(text)
    except Exception:
        return None


def _dedupe(urls: Iterable[str]) -> List[str]:
    seen = set()
    out = []
    for url in urls:
        if url and url not in seen:
            seen.add(url)
            out.append(url)
    return out
